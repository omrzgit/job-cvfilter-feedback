import streamlit as st
from openai import OpenAI
from fpdf import FPDF
import os
import io
import tempfile
import re

# ==========================================
# Page Configuration & Custom CSS Styling
# ==========================================
st.set_page_config(
    page_title="JOBIEE | AI Job Filter & Mock Interviewer",
    page_icon="💼",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom Styling for sleek modern UI
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&display=swap');

    html, body, [class*="css"] {
        font-family: 'Plus Jakarta Sans', sans-serif;
    }

    /* Gradient Headers */
    .main-title {
        font-size: 2.2rem;
        font-weight: 800;
        background: linear-gradient(135deg, #4F46E5 0%, #7C3AED 50%, #EC4899 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        margin-bottom: 0.2rem;
    }
    .sub-title {
        color: #64748B;
        font-size: 1rem;
        font-weight: 500;
        margin-bottom: 1.5rem;
    }

    /* Stepper Bar */
    .stepper-container {
        display: flex;
        justify-content: space-between;
        align-items: center;
        background: #F8FAFC;
        border: 1px solid #E2E8F0;
        border-radius: 16px;
        padding: 12px 20px;
        margin-bottom: 2rem;
    }
    @media (prefers-color-scheme: dark) {
        .stepper-container {
            background: #1E293B;
            border-color: #334155;
        }
    }
    .step-item {
        display: flex;
        align-items: center;
        gap: 8px;
        font-size: 0.88rem;
        font-weight: 600;
        color: #94A3B8;
    }
    .step-item.active {
        color: #4F46E5;
    }
    .step-item.completed {
        color: #10B981;
    }
    .step-badge {
        width: 26px;
        height: 26px;
        border-radius: 50%;
        display: flex;
        align-items: center;
        justify-content: center;
        font-size: 0.75rem;
        font-weight: 700;
        background: #E2E8F0;
        color: #64748B;
    }
    .step-item.active .step-badge {
        background: #4F46E5;
        color: #FFFFFF;
        box-shadow: 0 0 10px rgba(79, 70, 229, 0.4);
    }
    .step-item.completed .step-badge {
        background: #10B981;
        color: #FFFFFF;
    }

    /* Cards */
    .custom-card {
        background: #FFFFFF;
        border: 1px solid #E2E8F0;
        border-radius: 16px;
        padding: 24px;
        box-shadow: 0 4px 20px -2px rgba(0, 0, 0, 0.05);
        margin-bottom: 1.5rem;
    }
    @media (prefers-color-scheme: dark) {
        .custom-card {
            background: #0F172A;
            border-color: #1E293B;
        }
    }

    /* Highlight Card */
    .company-card {
        border-left: 5px solid #4F46E5;
        background: linear-gradient(135deg, rgba(79, 70, 229, 0.04) 0%, rgba(124, 58, 237, 0.02) 100%);
        border-radius: 12px;
        padding: 20px;
        margin-bottom: 1rem;
        border: 1px solid #E0E7FF;
    }
    @media (prefers-color-scheme: dark) {
        .company-card {
            background: rgba(79, 70, 229, 0.1);
            border-color: #3730A3;
        }
    }

    /* Pill Badges */
    .badge {
        display: inline-block;
        padding: 4px 12px;
        border-radius: 9999px;
        font-size: 0.75rem;
        font-weight: 700;
        text-transform: uppercase;
        letter-spacing: 0.05em;
    }
    .badge-primary { background: #EEF2FF; color: #4F46E5; }
    .badge-success { background: #ECFDF5; color: #059669; }
    .badge-warning { background: #FFFBEB; color: #D97706; }
    .badge-info { background: #F0F9FF; color: #0284C7; }

    /* Score meter */
    .score-circle {
        font-size: 2.5rem;
        font-weight: 800;
        color: #4F46E5;
        text-align: center;
    }
</style>
""", unsafe_allow_html=True)

# ==========================================
# Initialize Session State
# ==========================================
default_state = {
    'page': 0,
    'name': '',
    'degree': 'Graduate',
    'job_position': 'Junior',
    'education': 'Bachelor AI Engineering',
    'selected_company': '',
    'cv_text': '',
    'cv_filename': '',
    'questions': [],
    'answers': [],
    'evaluations': [],
    'current_question_index': 0,
    'openai_api_key': os.getenv("OPENAI_API_KEY", ""),
    'selected_model': "gpt-3.5-turbo"
}

for key, val in default_state.items():
    if key not in st.session_state:
        st.session_state[key] = val

# ==========================================
# Database / Companies Dataset
# ==========================================
companies_data = {
    'Undergraduate': {
        'Junior': [
            {
                'name': 'Motive Technologies',
                'Education': 'bachelor ai engineering',
                'Degree': 'Undergraduate',
                'position': 'Junior',
                'industry': 'Artificial Intelligence & IoT',
                'company_information': "A global technology company building intelligent hardware and AI software to revolutionize the safety and productivity of physical operations.",
                'job_description': 'Assist in developing, testing, and fine-tuning AI and machine learning models. Candidate must have strong programming skills in Python, familiarity with PyTorch/TensorFlow, and basic data processing techniques.'
            },
            {
                'name': 'Adept Tech Solutions',
                'Education': 'bachelor ai engineering',
                'Degree': 'Undergraduate',
                'position': 'Junior',
                'industry': 'Software & AI Services',
                'company_information': "An emerging enterprise software engineering agency focused on building AI automations, conversational agents, and data pipeline solutions.",
                'job_description': 'Support AI development projects by building and testing predictive models. Requires knowledge of Python, ML algorithms, Git version control, and data preprocessing.'
            },
            {
                'name': 'Coca Cola Beverages',
                'Education': 'bachelor business administration',
                'Degree': 'Undergraduate',
                'position': 'Internship',
                'industry': 'Consumer Goods & Retail',
                'company_information': "A world leader in the beverage sector, known for global distribution networks, marketing excellence, and strategic operations.",
                'job_description': 'Gain hands-on corporate experience in operational planning, supply coordination, and market research. Candidate should possess strong communication skills and analytical mindset.'
            }
        ],
        'Senior': []
    },
    'Graduate': {
        'Junior': [
            {
                'name': 'Jazz Digital Services',
                'Education': 'bachelor ai engineering',
                'Degree': 'Graduate',
                'position': 'Junior',
                'industry': 'Telecom & FinTech',
                'company_information': "The leading digital communications company in Pakistan offering nationwide connectivity, cloud computing, and digital financial ecosystems.",
                'job_description': 'Build, optimize, and deploy AI models for customer churn analytics, conversational chatbots, and network optimization. Proficiency in Python, SQL, and cloud basics required.'
            },
            {
                'name': 'Systems Limited',
                'Education': 'bachelor data science',
                'Degree': 'Graduate',
                'position': 'Junior',
                'industry': 'Enterprise IT & Consulting',
                'company_information': "A premier global IT solutions provider delivering cutting-edge enterprise software, cloud transformation, and data intelligence services.",
                'job_description': 'Analyze large-scale transactional and operational datasets. Create interactive PowerBI/Tableau dashboards and implement statistical models using Python and SQL.'
            }
        ],
        'Senior': [
            {
                'name': 'Systems Limited (Data Insights Group)',
                'Education': 'bachelor data science',
                'Degree': 'Graduate',
                'position': 'Senior',
                'industry': 'Big Data & Cloud Analytics',
                'company_information': "A specialized analytics division engineering modern data platforms, data warehouses, and AI-driven business intelligence for global Fortune 500 clients.",
                'job_description': 'Architect end-to-end data analytics pipelines, design predictive frameworks, and lead junior data engineers. Candidate must possess advanced mastery in Python, R, SQL, and cloud platforms.'
            },
            {
                'name': 'Increative Digital Agency',
                'Education': 'bachelor website development',
                'Degree': 'Graduate',
                'position': 'Senior',
                'industry': 'Web & Digital Experience',
                'company_information': "A boutique full-stack digital product studio crafting high-converting modern web applications, progressive web apps (PWAs), and e-commerce platforms.",
                'job_description': 'Lead full-stack web development initiatives, ensuring maximum speed, scalability, and UX fidelity. Proficiency in JavaScript/TypeScript, React/Next.js, Node.js, and CSS architectures required.'
            },
            {
                'name': 'Engro Corporation',
                'Education': 'bachelor business administration',
                'Degree': 'Graduate',
                'position': 'Senior',
                'industry': 'Conglomerate & Energy',
                'company_information': "One of Pakistan's most respected conglomerates with diversified operations across fertilizers, petrochemicals, energy infrastructure, and digital logistics.",
                'job_description': 'Drive cross-functional strategic projects, manage department budgets, optimize supply chain pipelines, and mentor junior management associates.'
            }
        ]
    },
    'Post Graduate': {
        'Junior': [],
        'Senior': [
            {
                'name': 'Coding Souls Tech Lab',
                'Education': 'bachelor software engineering',
                'Degree': 'Post Graduate',
                'position': 'Senior',
                'industry': 'Custom Software & Cloud Systems',
                'company_information': "An engineering-first software studio developing microservices architectures, high-concurrency cloud applications, and distributed SaaS products.",
                'job_description': 'Lead senior engineering teams, design scalable distributed software architectures, conduct code reviews, and drive system reliability. Deep knowledge in Python, Go/C++, and CI/CD pipelines.'
            }
        ]
    }
}

# ==========================================
# Helper Functions (CV Parsing, OpenAI, PDF)
# ==========================================
def extract_text_from_file(uploaded_file):
    """Extract text from uploaded CV file (PDF, DOCX, TXT)"""
    if uploaded_file is None:
        return ""
    try:
        filename = uploaded_file.name.lower()
        if filename.endswith(".txt"):
            return uploaded_file.read().decode("utf-8", errors="ignore")
        elif filename.endswith(".pdf"):
            try:
                import pypdf
                reader = pypdf.PdfReader(uploaded_file)
                text = "\n".join([page.extract_text() or "" for page in reader.pages])
                return text[:4000] # trim reasonable length
            except Exception:
                return f"[PDF Resume File: {uploaded_file.name}]"
        elif filename.endswith(".docx"):
            try:
                import docx
                doc = docx.Document(uploaded_file)
                text = "\n".join([p.text for p in doc.paragraphs])
                return text[:4000]
            except Exception:
                return f"[Word Document Resume: {uploaded_file.name}]"
        else:
            return f"[Uploaded CV: {uploaded_file.name}]"
    except Exception as e:
        return f"[File text extraction note: {str(e)}]"

def get_openai_client():
    """Returns OpenAI client if key is configured"""
    api_key = st.session_state.openai_api_key or os.getenv("OPENAI_API_KEY", "")
    if not api_key:
        return None
    return OpenAI(api_key=api_key)

def sanitize_for_pdf(text):
    """Sanitizes text to safe ASCII/Latin-1 compatible strings for FPDF"""
    if not text:
        return ""
    replacements = {
        '\u2018': "'", '\u2019': "'",
        '\u201c': '"', '\u201d': '"',
        '\u2013': '-', '\u2014': '--',
        '\u2022': '*', '\u2026': '...',
        '\u00a0': ' ', '\t': '    '
    }
    for orig, rep in replacements.items():
        text = text.replace(orig, rep)
    # Filter out remaining unsupported unicode characters
    return text.encode('latin-1', 'replace').decode('latin-1')

def render_stepper():
    """Renders the horizontal progress stepper"""
    steps = [
        ("Profile", 0),
        ("Match Role", 1),
        ("Role & CV", 2),
        ("AI Interview", 3),
        ("Scorecard", 4)
    ]
    current = st.session_state.page
    
    html = '<div class="stepper-container">'
    for label, idx in steps:
        if idx < current:
            status = "completed"
            icon = "✓"
        elif idx == current:
            status = "active"
            icon = str(idx + 1)
        else:
            status = ""
            icon = str(idx + 1)
        
        html += f'''
        <div class="step-item {status}">
            <div class="step-badge">{icon}</div>
            <span>{label}</span>
        </div>
        '''
        if idx < len(steps) - 1:
            html += '<div style="flex: 1; height: 2px; background: #E2E8F0; margin: 0 10px;"></div>'
    html += '</div>'
    st.markdown(html, unsafe_allow_html=True)

# ==========================================
# Sidebar Controls
# ==========================================
with st.sidebar:
    st.markdown("### 💼 **JOBIEE Navigator**")
    st.caption("AI-Powered Job Match & Interview Simulator")
    st.divider()

    # Candidate Profile Summary
    st.markdown("#### 👤 **Applicant Profile**")
    if st.session_state.name:
        st.write(f"**Name:** {st.session_state.name}")
        st.write(f"**Degree:** {st.session_state.degree}")
        st.write(f"**Position:** {st.session_state.job_position}")
        st.write(f"**Field:** {st.session_state.education}")
    else:
        st.info("Complete step 1 to build your candidate profile.")

    if st.session_state.selected_company:
        st.divider()
        st.markdown("#### 🏢 **Target Company**")
        st.success(f"**{st.session_state.selected_company}**")

    st.divider()
    # OpenAI Settings
    st.markdown("#### ⚙️ **AI Engine Settings**")
    api_key_input = st.text_input(
        "OpenAI API Key",
        value=st.session_state.openai_api_key,
        type="password",
        help="Enter your OpenAI API key or configure OPENAI_API_KEY environment variable."
    )
    if api_key_input != st.session_state.openai_api_key:
        st.session_state.openai_api_key = api_key_input
        st.rerun()

    model_choice = st.selectbox(
        "AI Model",
        ["gpt-3.5-turbo", "gpt-4o-mini", "gpt-4o"],
        index=0
    )
    st.session_state.selected_model = model_choice

    if st.session_state.openai_api_key:
        st.success("API Key Active ✓", icon="🟢")
    else:
        st.warning("API Key needed for live AI generation.", icon="⚠️")

    st.divider()
    if st.button("🔄 Restart Process", use_container_width=True):
        for key in list(st.session_state.keys()):
            if key not in ['openai_api_key', 'selected_model']:
                st.session_state[key] = default_state[key]
        st.rerun()

# ==========================================
# View 1: Personal Information
# ==========================================
def personal_info():
    render_stepper()
    
    col1, col2 = st.columns([2, 1])
    with col1:
        st.markdown('<div class="main-title">Candidate Profile</div>', unsafe_allow_html=True)
        st.markdown('<div class="sub-title">Tell us about your educational background and target career level.</div>', unsafe_allow_html=True)

    with st.container():
        st.markdown('<div class="custom-card">', unsafe_allow_html=True)
        
        c1, c2 = st.columns(2)
        with c1:
            name = st.text_input("Full Name", value=st.session_state.name, placeholder="e.g. Alex Morgan")
            degree = st.selectbox("Current Degree Level", ["Undergraduate", "Graduate", "Post Graduate"], 
                                  index=["Undergraduate", "Graduate", "Post Graduate"].index(st.session_state.degree) if st.session_state.degree in ["Undergraduate", "Graduate", "Post Graduate"] else 1)
        with c2:
            job_position = st.selectbox("Target Seniority", ["Junior", "Senior"], 
                                        index=["Junior", "Senior"].index(st.session_state.job_position) if st.session_state.job_position in ["Junior", "Senior"] else 0)
            
            education_suggestions = [
                "Bachelor AI Engineering",
                "Bachelor Data Science",
                "Bachelor Software Engineering",
                "Bachelor Website Development",
                "Bachelor Business Administration"
            ]
            education = st.selectbox("Field of Study / Degree Title", education_suggestions)
        
        st.markdown("<br>", unsafe_allow_html=True)
        if st.button("🔍 Find Matching Companies ➔", type="primary", use_container_width=True):
            if name.strip() and education.strip():
                st.session_state.name = name.strip()
                st.session_state.degree = degree
                st.session_state.job_position = job_position
                st.session_state.education = education
                st.session_state.page = 1
                st.rerun()
            else:
                st.error("Please provide your full name before continuing.")
        
        st.markdown('</div>', unsafe_allow_html=True)

# ==========================================
# View 2: Company Selection
# ==========================================
def company_selection():
    render_stepper()
    
    st.markdown('<div class="main-title">Matched Opportunities</div>', unsafe_allow_html=True)
    st.markdown('<div class="sub-title">Based on your academic background and desired seniority level:</div>', unsafe_allow_html=True)

    degree = st.session_state.degree
    job_position = st.session_state.job_position
    education = st.session_state.education.lower()

    matched_companies = []
    for deg_key, pos_dict in companies_data.items():
        if deg_key.lower() == degree.lower():
            for pos_key, comp_list in pos_dict.items():
                if pos_key.lower() == job_position.lower():
                    for comp in comp_list:
                        if comp['Education'].lower() in education or education in comp['Education'].lower():
                            matched_companies.append(comp)

    if matched_companies:
        st.success(f"🎯 Found **{len(matched_companies)}** matching opening(s) for your profile!")
        
        for idx, comp in enumerate(matched_companies):
            with st.container():
                st.markdown(f'''
                <div class="company-card">
                    <div style="display:flex; justify-content:space-between; align-items:center; margin-bottom: 8px;">
                        <h3 style="margin:0; color:#1E293B;">{comp['name']}</h3>
                        <span class="badge badge-primary">{comp.get('industry', 'Technology')}</span>
                    </div>
                    <p style="color:#475569; font-size:0.95rem; margin-bottom:12px;">{comp['company_information']}</p>
                    <div style="display:flex; gap:10px;">
                        <span class="badge badge-info">🎓 {comp['Degree']}</span>
                        <span class="badge badge-success">💼 {comp['position']} Role</span>
                    </div>
                </div>
                ''', unsafe_allow_html=True)
                
                if st.button(f"Apply & Prepare for {comp['name']} ➔", key=f"btn_comp_{idx}", type="primary"):
                    st.session_state.selected_company = comp['name']
                    st.session_state.page = 2
                    st.rerun()

    else:
        st.warning("⚠️ No direct company match found for the exact criteria.")
        st.info("Try adjusting your degree level or selecting another field of study.")
        if st.button("⬅ Adjust Profile", use_container_width=True):
            st.session_state.page = 0
            st.rerun()

# ==========================================
# View 3: Company Details & CV Upload
# ==========================================
def company_info():
    render_stepper()
    
    selected = st.session_state.selected_company
    st.markdown(f'<div class="main-title">{selected}</div>', unsafe_allow_html=True)
    st.markdown('<div class="sub-title">Role details, requirements, and resume customization.</div>', unsafe_allow_html=True)

    company_obj = None
    for deg_key, pos_dict in companies_data.items():
        for pos_key, comp_list in pos_dict.items():
            for comp in comp_list:
                if comp['name'] == selected:
                    company_obj = comp
                    break

    col1, col2 = st.columns([3, 2], gap="large")
    
    with col1:
        st.markdown('<div class="custom-card">', unsafe_allow_html=True)
        st.markdown("### 🏢 **About the Company**")
        st.write(company_obj.get('company_information', 'Information not available.'))
        st.markdown("---")
        st.markdown("### 📋 **Job Description & Requirements**")
        st.write(company_obj.get('job_description', 'Description not available.'))
        st.markdown('</div>', unsafe_allow_html=True)

    with col2:
        st.markdown('<div class="custom-card">', unsafe_allow_html=True)
        st.markdown("### 📄 **Upload Resume / CV**")
        st.caption("Upload your CV to let our AI generate tailored questions based on your background.")
        
        cv_file = st.file_uploader("Upload PDF, DOCX, or TXT", type=["pdf", "docx", "txt"])
        
        if cv_file is not None:
            extracted = extract_text_from_file(cv_file)
            st.session_state.cv_text = extracted
            st.session_state.cv_filename = cv_file.name
            st.success(f"✓ Uploaded: **{cv_file.name}**")
            with st.expander("👁 Preview Extracted Resume Content"):
                st.write(extracted[:600] + "..." if len(extracted) > 600 else extracted)

        st.markdown("<br>", unsafe_allow_html=True)
        btn_col1, btn_col2 = st.columns(2)
        with btn_col1:
            if st.button("⬅ Back", use_container_width=True):
                st.session_state.page = 1
                st.rerun()
        with btn_col2:
            if st.button("Start Interview ➔", type="primary", use_container_width=True):
                st.session_state.page = 3
                st.session_state.questions = []
                st.session_state.answers = []
                st.session_state.evaluations = []
                st.session_state.current_question_index = 0
                st.rerun()
        st.markdown('</div>', unsafe_allow_html=True)

# ==========================================
# View 4: AI Mock Interview
# ==========================================
def interview_questions():
    render_stepper()
    
    client = get_openai_client()
    selected_company = st.session_state.selected_company
    job_position = st.session_state.job_position

    # Find job description
    job_description = "Software and Analytical Role"
    for deg_key, pos_dict in companies_data.items():
        for pos_key, comp_list in pos_dict.items():
            for comp in comp_list:
                if comp['name'] == selected_company:
                    job_description = comp.get('job_description', job_description)

    # 1. Generate questions if empty
    if not st.session_state.questions:
        if not client:
            st.error("🔑 OpenAI API Key required to generate dynamic interview questions.")
            st.info("Enter your OpenAI API key in the sidebar to begin the interview.")
            return

        with st.spinner("🤖 AI Coach is analyzing the role & CV to generate interview questions..."):
            is_coding = any(k in job_description.lower() for k in ['coding', 'python', 'java', 'c++', 'developer', 'software', 'model'])
            difficulty = "foundational/entry-level" if job_position.lower() == "junior" else "advanced/senior-level"
            
            prompt = f"""
            You are an expert technical interviewer for {selected_company}.
            Job Description: {job_description}
            Target Position: {job_position} ({difficulty})
            Candidate Resume text: {st.session_state.cv_text or "No CV provided"}

            Generate exactly 4 distinct, highly relevant interview questions tailored to this role and candidate:
            - If it's a technical role, include 2 conceptual/coding problem questions and 2 behavioral/system design questions.
            - If non-technical, generate 4 role-specific situational and strategic questions.
            
            Format: Output only the 4 questions, one per line, numbered 1 to 4. Do not include markdown preamble.
            """
            
            try:
                response = client.chat.completions.create(
                    model=st.session_state.selected_model,
                    messages=[
                        {"role": "system", "content": "You are a professional hiring manager conducting a structured interview."},
                        {"role": "user", "content": prompt}
                    ],
                    temperature=0.7
                )
                raw_lines = response.choices[0].message.content.strip().split("\n")
                cleaned_questions = [re.sub(r'^\d+[\.\)]\s*', '', q.strip()) for q in raw_lines if q.strip()]
                st.session_state.questions = cleaned_questions[:4]
            except Exception as e:
                st.error(f"Error calling OpenAI API: {str(e)}")
                # Fallback questions
                st.session_state.questions = [
                    f"Why are you interested in joining {selected_company} for the {job_position} position?",
                    "Can you describe a challenging project you worked on and how you overcame key obstacles?",
                    "How do your skills and background align with the core requirements of this role?",
                    "Where do you see yourself technically and professionally in the next 2-3 years?"
                ]

    questions = st.session_state.questions
    total_q = len(questions)
    curr_idx = st.session_state.current_question_index

    # Progress Bar
    progress_val = min(1.0, (curr_idx) / total_q)
    st.progress(progress_val)
    st.caption(f"Question **{min(curr_idx + 1, total_q)}** of **{total_q}**")

    if curr_idx < total_q:
        current_q = questions[curr_idx]
        
        st.markdown(f'''
        <div class="custom-card">
            <span class="badge badge-primary">Question {curr_idx + 1}</span>
            <h3 style="margin-top: 10px; color: #1E293B;">{current_q}</h3>
        </div>
        ''', unsafe_allow_html=True)

        user_answer = st.text_area(
            "Your Answer:",
            height=160,
            placeholder="Type your response here. Be specific, articulate, and include examples where applicable...",
            key=f"ans_input_{curr_idx}"
        )

        col_btn1, col_btn2 = st.columns([4, 1])
        with col_btn2:
            submit_btn = st.button("Submit & Evaluate ➔", type="primary", use_container_width=True)

        if submit_btn:
            if not user_answer.strip():
                st.warning("Please provide an answer before moving forward.")
            else:
                score = 7
                feedback = "Good response covering the essential points."
                
                if client:
                    with st.spinner("🧠 AI is analyzing and scoring your response..."):
                        eval_prompt = f"""
                        Question: {current_q}
                        Candidate's Answer: {user_answer}
                        Role: {job_position} at {selected_company}

                        Evaluate the answer for accuracy, depth, clarity, and communication effectiveness.
                        Start your response with:
                        Score: X/10 (where X is an integer 1-10)
                        Followed by a concise 2-3 sentence constructive feedback explaining strengths and areas for improvement.
                        """
                        try:
                            eval_res = client.chat.completions.create(
                                model=st.session_state.selected_model,
                                messages=[
                                    {"role": "system", "content": "You are a constructive and insightful hiring evaluator."},
                                    {"role": "user", "content": eval_prompt}
                                ],
                                temperature=0.3
                            )
                            feedback_text = eval_res.choices[0].message.content.strip()
                            # Parse score
                            score_match = re.search(r'Score:\s*(\d+)', feedback_text, re.IGNORECASE)
                            if score_match:
                                score = int(score_match.group(1))
                            feedback = feedback_text
                        except Exception as e:
                            feedback = f"Score: 7/10\nEvaluation saved. (API Notice: {str(e)})"

                # Store evaluation & answer
                st.session_state.answers.append({
                    'question': current_q,
                    'answer': user_answer,
                    'score': score,
                    'feedback': feedback
                })
                
                st.session_state.current_question_index += 1
                if st.session_state.current_question_index >= total_q:
                    st.session_state.page = 4
                st.rerun()

    else:
        st.session_state.page = 4
        st.rerun()

# ==========================================
# View 5: Scorecard & PDF Report
# ==========================================
def generate_report():
    render_stepper()
    
    st.markdown('<div class="main-title">🎉 Interview Performance Scorecard</div>', unsafe_allow_html=True)
    st.markdown('<div class="sub-title">Detailed breakdown of your responses, AI scores, and downloadable PDF report.</div>', unsafe_allow_html=True)

    answers = st.session_state.get('answers', [])
    total_answers = len(answers)
    
    if total_answers > 0:
        avg_score = sum(item['score'] for item in answers) / total_answers
    else:
        avg_score = 0

    # Metric Row
    m1, m2, m3, m4 = st.columns(4)
    with m1:
        st.markdown(f'''
        <div class="custom-card" style="text-align:center;">
            <div class="score-circle">{avg_score:.1f}<span style="font-size:1.2rem; color:#94A3B8;">/10</span></div>
            <div style="font-weight:700; color:#475569;">Overall Score</div>
        </div>
        ''', unsafe_allow_html=True)
    with m2:
        performance_tag = "Outstanding" if avg_score >= 8.5 else "Strong Candidate" if avg_score >= 7.0 else "Needs Practice"
        badge_type = "badge-success" if avg_score >= 7.0 else "badge-warning"
        st.markdown(f'''
        <div class="custom-card" style="text-align:center;">
            <div style="margin-top:10px;"><span class="badge {badge_type}" style="font-size:1rem; padding:8px 16px;">{performance_tag}</span></div>
            <div style="font-weight:700; color:#475569; margin-top:16px;">Readiness Rating</div>
        </div>
        ''', unsafe_allow_html=True)
    with m3:
        st.markdown(f'''
        <div class="custom-card" style="text-align:center;">
            <div class="score-circle" style="color:#0284C7;">{total_answers}</div>
            <div style="font-weight:700; color:#475569;">Questions Completed</div>
        </div>
        ''', unsafe_allow_html=True)
    with m4:
        st.markdown(f'''
        <div class="custom-card" style="text-align:center;">
            <div style="font-size:1.2rem; font-weight:800; color:#1E293B; margin-top:10px;">{st.session_state.selected_company}</div>
            <div style="font-weight:700; color:#475569; margin-top:14px;">Target Organization</div>
        </div>
        ''', unsafe_allow_html=True)

    # Detailed Q&A Accordion
    st.markdown("### 📝 **Question-by-Question Breakdown**")
    for idx, item in enumerate(answers):
        with st.expander(f"Q{idx+1}: {item['question']} — (Score: {item['score']}/10)", expanded=(idx == 0)):
            st.markdown(f"**Your Answer:**")
            st.info(item['answer'])
            st.markdown(f"**AI Evaluation & Insights:**")
            st.success(item['feedback'])

    # Build PDF Document
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    # Title Banner
    pdf.set_font("Arial", 'B', 18)
    pdf.cell(0, 12, txt=sanitize_for_pdf("JOBIEE - AI Interview & Candidate Assessment"), ln=True, align="C")
    pdf.ln(5)

    # Candidate Meta
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(0, 8, txt=sanitize_for_pdf("Candidate Summary"), ln=True)
    pdf.set_font("Arial", size=10)
    pdf.cell(0, 6, txt=sanitize_for_pdf(f"Candidate Name: {st.session_state.name or 'N/A'}"), ln=True)
    pdf.cell(0, 6, txt=sanitize_for_pdf(f"Degree Level: {st.session_state.degree} | Field: {st.session_state.education}"), ln=True)
    pdf.cell(0, 6, txt=sanitize_for_pdf(f"Target Role: {st.session_state.job_position} @ {st.session_state.selected_company}"), ln=True)
    pdf.cell(0, 6, txt=sanitize_for_pdf(f"Overall Benchmark Score: {avg_score:.2f} / 10.0 ({performance_tag})"), ln=True)
    pdf.ln(6)

    # Questions & Answers
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(0, 8, txt=sanitize_for_pdf("Interview Transcript & Evaluation"), ln=True)
    pdf.ln(2)

    for idx, item in enumerate(answers):
        pdf.set_font("Arial", 'B', 10)
        pdf.multi_cell(0, 6, txt=sanitize_for_pdf(f"Q{idx+1}: {item['question']} [Score: {item['score']}/10]"))
        pdf.set_font("Arial", size=9)
        pdf.multi_cell(0, 5, txt=sanitize_for_pdf(f"Candidate Answer: {item['answer']}"))
        pdf.set_font("Arial", 'I', 9)
        pdf.multi_cell(0, 5, txt=sanitize_for_pdf(f"AI Feedback: {item['feedback']}"))
        pdf.ln(4)

    # Output to buffer
    pdf_bytes = io.BytesIO()
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        pdf.output(tmp.name)
        tmp.seek(0)
        pdf_data = tmp.read()
        pdf_bytes = io.BytesIO(pdf_data)

    st.divider()
    c_down1, c_down2 = st.columns([2, 1])
    with c_down1:
        st.download_button(
            label="📥 Download Official Assessment Report (PDF)",
            data=pdf_bytes,
            file_name=f"JOBIEE_Interview_Report_{st.session_state.name.replace(' ', '_') or 'Candidate'}.pdf",
            mime="application/pdf",
            type="primary",
            use_container_width=True
        )
    with c_down2:
        if st.button("🔁 Start New Interview", use_container_width=True):
            for key in list(st.session_state.keys()):
                if key not in ['openai_api_key', 'selected_model']:
                    st.session_state[key] = default_state[key]
            st.rerun()

# ==========================================
# Main Router
# ==========================================
def main():
    page = st.session_state.page
    if page == 0:
        personal_info()
    elif page == 1:
        company_selection()
    elif page == 2:
        company_info()
    elif page == 3:
        interview_questions()
    elif page == 4:
        generate_report()

if __name__ == "__main__":
    main()